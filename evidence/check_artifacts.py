"""Independent file-member and document-content checks for the frozen audit inputs."""
from pathlib import Path
import json,hashlib,re,tarfile,zipfile,unicodedata,sys,subprocess
from docx import Document
from PIL import Image
import fitz
A=Path(__file__).resolve().parent.parent;B=A/'evidence';S=A/'work/PACEgame_v0_2_11';P=A/'manuscript';C=A/'outputs'
load=lambda p:json.loads(p.read_text());sha=lambda b:hashlib.sha256(b).hexdigest()
def eligible(p):return p.is_file() and '__pycache__' not in p.parts and p.suffix!='.pyc' and not p.name.startswith('.codex-library')
source={p.relative_to(S).as_posix():sha(p.read_bytes()) for p in (S/'pacegame').rglob('*') if eligible(p)}
with zipfile.ZipFile(S/'dist/pacegame-0.2.11-py3-none-any.whl') as z:
 assert z.testzip() is None
 wheel={n:sha(z.read(n)) for n in z.namelist() if n.startswith('pacegame/') and not n.endswith('/')}
with tarfile.open(S/'dist/pacegame-0.2.11.tar.gz') as t:
 sdist={'/'.join(m.name.split('/')[1:]):sha(t.extractfile(m).read()) for m in t.getmembers() if m.isfile() and '/pacegame/' in m.name}
assert source==wheel==sdist and len(source)==61
installed=[]
for name in ('wheel_env','core_env'):
 paths=list((A/name).glob('lib/python*/site-packages/pacegame'));assert len(paths)==1
 got={'pacegame/'+p.relative_to(paths[0]).as_posix():sha(p.read_bytes()) for p in paths[0].rglob('*') if eligible(p)}
 assert source==got;installed.append({'environment':name,'package_members':len(got),'byte_equal':True})
for p in (S/'pacegame/examples').iterdir():
 if eligible(p):assert p.read_bytes()==(S/'examples'/p.name).read_bytes()
doc=Document(C/'PACEgame_SoftwareX_Manuscript_v0_2_11.docx');paras=[p.text for p in doc.paragraphs]
c=load(P/'manuscript_content.json');tex=(P/'manuscript.tex').read_text()
def flatten(xs):
 for x in xs:
  if isinstance(x,str):yield x
  elif x[0]=='code':yield x[1]
  else:yield from flatten(x[1])
parts=[s for _,xs in c['sections'] for s in flatten(xs)];cursor=0
for s in parts:cursor=paras.index(s,cursor)+1
wc=lambda s:len(re.findall(r"\b[\w]+(?:[’'-][\w]+)*\b",s))
assert wc(' '.join(parts))==2979 and wc(c['abstract'])==248
def escape(s):
 maps={'^':r'\textasciicircum{}','~':r'\textasciitilde{}','\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','×':r'$\times$','≤':r'$\leq$','−':r'$-$'}
 for old,new in [('⁻¹⁶','\ue000'),('⁻⁶','\ue001'),('⁻¹⁷','\ue002')]:s=s.replace(old,new)
 maps.update({'\ue000':r'$^{-16}$','\ue001':r'$^{-6}$','\ue002':r'$^{-17}$'})
 return ''.join(maps.get(x,x) for x in s)
for s in parts:assert (s if '\n' in s else escape(s)) in tex
cells=[cell.text for t in doc.tables for row in t.rows for cell in row.cells];assert len(cells)==232
assert all(escape(s) in tex for s in cells)
for title,s in c['declarations']:assert title in paras and s in paras and escape(s) in tex
for s in c['references']:
 assert s in paras
 for chunk in re.split(r'(https?://\S+)',s):assert (r'\url{'+chunk.rstrip('.')+'}' if chunk.startswith('http') else escape(chunk)) in tex
metadata=load(P/'metadata_tables.json');rows=[[tuple(cell.text for cell in row.cells) for row in t.rows[1:]] for t in doc.tables[:2]]
assert rows[0]==[tuple(x) for x in metadata['code']] and rows[1]==[tuple(x) for x in metadata['software']]
expected=load(B/'depth_report_fields.json')['fields'];assert [[r[0],r[1]] for rs in rows for r in rs]==expected and len(expected)==17
assert metadata==load(S/'evidence/metadata_tables.json')
assert rows[0][0][2]==rows[1][0][2]=='0.2.11'
assert (P/'PACEgame_SoftwareX_Manuscript_v0_2_11.docx').read_bytes()==(C/'PACEgame_SoftwareX_Manuscript_v0_2_11.docx').read_bytes()
norm=lambda s:''.join(x for x in unicodedata.normalize('NFKC',s) if x.isalnum())
def pdftext(p):
 chunks=[]
 with fitz.open(p) as f:
  for page in f:
   for b in page.get_text('blocks'):
    assert b[0]>=-.5 and b[1]>=-.5 and b[2]<=page.rect.width+.5 and b[3]<=page.rect.height+.5
    if not (b[4].strip()==str(page.number+1) and b[1]>page.rect.height*.9):chunks.append(b[4])
  return norm(' '.join(chunks)),len(f)
wordtext,np=pdftext(B/'manuscript_render/PACEgame_SoftwareX_Manuscript_v0_2_11.pdf');_,nt=pdftext(B/'tex_build/manuscript.pdf')
texts=[s for s in paras+cells if s.strip()];assert all(norm(s) in wordtext for s in texts) and np>0 and nt>0
assert len(doc.tables)==5 and len(doc.inline_shapes)==3
for p in (P/'figures').glob('*.png'):
 im=Image.open(p).convert('RGB');assert all(r==g==b for r,g,b in im.getdata())
assert all(len(s)<=85 for s in c['highlights'])
data={'version':'0.2.11','package_members':61,'source_wheel_sdist_byte_equal':True,'installed_package_checks':installed,'examples_mirrored':True,
 'body_words':2979,'abstract_words':248,'body_parts_checked':len(parts),'table_cells_checked':len(cells),'metadata_fields_matched':17,'word_tex_body_and_table_cells_equal':True,'all_nonempty_word_texts_rendered':True,'nonempty_word_texts_checked':len(texts),'word_pages':np,'tex_pages':nt,'tables':5,'figures':3,'all_figures_grayscale':True,'highlights_characters':[len(s) for s in c['highlights']],
 'official_template_certified':False,'fresh_sdist_rebuild_performed':True,'rebuilt_wheel_bytes_identical':load(B/'distribution_integrity.json')['wheel_bytes_identical']}
(B/'artifact_validation.json').write_text(json.dumps(data,indent=2)+'\n');print(json.dumps(data,indent=2))
