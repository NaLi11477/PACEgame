from pathlib import Path
import json,re,shutil,os
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
R=Path(__file__).resolve().parent
c=json.loads((R/'manuscript_content.json').read_text())
validation_inputs=Path(os.environ.get('PACEGAME_VALIDATION_DIR',str(R/'validation_source')))
r=json.loads((validation_inputs/'release_validation.json').read_text())
a=json.loads((validation_inputs/'audit_021_acceptance.json').read_text())
a3=json.loads((validation_inputs/'audit_022_acceptance.json').read_text())
a5=json.loads((validation_inputs/'audit_024_acceptance.json').read_text())
a4=json.loads((validation_inputs/'audit_023_acceptance.json').read_text())
root=R.parent/'outputs';root.mkdir(exist_ok=True)
pkg=R.parent/'manuscript';pkg.mkdir(exist_ok=True)
figdir=pkg/'figures';figdir.mkdir(exist_ok=True)
metadata=json.loads((pkg/'metadata_tables.json').read_text())
metadata1=[tuple(x) for x in metadata['code']];metadata2=[tuple(x) for x in metadata['software']]
comparison=[('Tool','Primary stated scope','Relationship to this workflow'),('Gambit [1]','Finite extensive and strategic games','A full game-analysis toolkit; continuous-time symbolic integration is not this comparison’s input representation.'),('Nashpy [2]','Two-player normal-form equilibria','Useful for payoff-matrix games; PACEgame inspects conditional timing scenarios.'),('QuantEcon [3]','Computational economics library','Broad reusable numerical routines; domain-specific inspection is supplied here.'),('SymPy [4]','Symbolic mathematics','Provides the algebraic foundation used by PACEgame.'),('PACEgame 0.2.11','Symbolic timing models and sampled inspection','Combines primitives, bounded candidate search, local diagnostics and claim coverage; no Nash/global certificate.')]
numerics=[('Configuration','RL signal / entry','RL payoff','DF / DL payoff','RL state')]
for row in r['cases']:
 sol=row['solutions']['RL'];d=sol['point']['delta'];e=sol['point']['varsigma']
 numerics.append((row['label'],f'{d:.6g} / {e:.6g}' if sol['rankable'] else 'search budget / 0','%.9g'%sol['payoffs']['x'] if sol['rankable'] else 'excluded',f"{row['solutions']['DF']['payoffs']['x']:.9g} / {row['solutions']['DL']['payoffs']['x']:.9g}",sol['kind']['delta']+'; '+sol['kind']['varsigma']))
figures=[('fig1_architecture','Fig. 1. Model specification, numerical search, diagnostics and review artifacts. Every search result retains its declared limitations.'),('fig2_payoffs','Fig. 2. Focal net payoffs at the illustrative point for the stock-cost, bounded-signal and readiness-lag variants. Values come from release_validation.json.'),('fig3_claims','Fig. 3. The same six claims over nine parameter samples per configuration. P: passed; F: failed; S: skipped. No computation errors occurred in these runs.')]
g=a['groups']
edge_rows=[('Acceptance case','Input / scope','Recorded result'),
 ('Sweep status','unresolved / complete / malformed',' / '.join(str(x['exit']) for x in g['R2-01']['commands'])),
 ('Analytic crossings','phi=1; scans 2, 3, 12','both APIs recover 1'),
 ('Fixed entry','stock-cost; entry=0.3','signal %.6g; payoff %.8g'%(g['R2-03']['solution']['point']['delta'],g['R2-03']['solution']['payoffs']['x'])),
 ('Unknown sensitivity','quality y=1+floor(phi)','Competitive skipped 3/3'),
 ('Short horizon','Logistic; horizon=0.5','entry %.6g; rankable'%g['R2-05']['short_logistic']['point']['varsigma']),
 ('Long horizon','TwoPhase; horizon=4, entry=3','payoff %.6g; ceiling %.6g'%(g['R2-05']['long_two_phase']['payoffs']['x'],g['R2-05']['long_ceiling'])),
 ('Required checks','unknown bounds / omitted gross','incomplete / negative net allowed'),
 ('Interiority scope','unresolved RF parent','skipped 1/1; incomplete report')]

g3=a3['groups']
edge_rows += [
 ('Slotted public state','gain 1 to 2; silent entry', '%.7g to %.7g'%(g3['R3-01']['before']['payoffs']['x'],g3['R3-01']['after_public_mutation']['payoffs']['x'])),
 ('Cost-only response','timing_cost=phi; shares=1/2','Competitive failed %d/3'%g3['R3-02']['confounded_report']['results'][0]['failed']),
 ('Claim iterables','list / tuple / iterator / generator','each fails 3/3; complete accounting'),
 ('Invalid tolerance','NaN / infinity / negative / complex','all rejected before checks'),
 ('Difference overflow','finite extreme payoffs','Sign: errors %d; passed %d'%(g3['R3-05']['sign_report']['results'][0]['errors'],g3['R3-05']['sign_report']['results'][0]['passed']))]

g4=a4['groups']
edge_rows += [
 ('Typed dictionary keys','integer 1 to string 1','payoff 51/8000 to 51/16000; fresh/invalidate agree'),
 ('Bound payoff targets','fixed entry .1; x / y','phi slopes +160/16136289 and negative; rival > fails, < passes'),
 ('Invalid Sign request','unknown target or malformed expression','seven requests reject before sample evaluation'),
 ('One-shot strategy pools','five iterable forms; three phi points','identical scans; invalid pools reject before solves')]

edge_rows += [
 ('Candidate metadata','wrong kind / missing coordinate','rejected; true upper and singleton pass'),
 ('Actual player quality','unknown z / partial x=2','unknown rejects; default y=1 agrees; no role gap'),
 ('Release CI gate','all packaged tests and exporters','493 tests discovered; 41 audit groups configured'),
 ('Expression coefficients','17/20 and 2*alpha','no external phi / free alpha; equivalent payoffs'),
 ('Metadata labels','depth report C1-C9 and S1-S8','17/17 exact; support email pending')
]

edge_rows += [
 ('Interval precision','A=10^20; [A,A+1]','typed rejection; singleton and rescaled controls pass'),
 ('Structural symbols','rate/time and entry/signal aliases','construction and mutable boundaries reject'),
 ('Critical boundary curvature','audited joint corner; infeasible eigenvector','positive feasible curvature fails; infeasible direction passes'),
 ('Distinct player identities','x/x versus alpha/beta','duplicates reject; two payoff keys retained'),
 ('Reproduction dependencies','core and optional principal pins','Matplotlib 3.10.8; NumPy 2.3.5; optional lock recorded')
]

edge_rows += [
 ('Retained candidate precision','A=10^20; [A+1,A+32769]','typed rejection; shifted lawful net=-1'),
 ('Fixed domain and payoff precision','[A+1,A+1]; cancellation control','unsafe singleton rejects; payoff disagreement excludes'),
 ('Remaining scalar boundary','zero gradient; fixed neighbor','positive curvature fails; concave and strict pass')
]

edge_rows += [
 ('Parameter objective retention','A=10^20 or 2^60; linear / squared','point=1; payoff=1; wrong lower fails'),
 ('Candidate expression retention','offset-A+1 and literal 1; target -u²','identical upper derivative -2; both fail'),
 ('Necessary derivative retention','reordered symbols; interior / mixed Hessian','correct upper passes; positive critical curvature fails')
]

edge_rows += [
 ('Retained joint cone precision','A=2^53 or 2^60; exact eigenvalue 2','wrong origin fails; unstable expanded solve excluded'),
 ('Shared interior Hessian','automatic / explicit metadata; eigenvalues 1,-5','saddle fails; concave and main-solver controls pass'),
 ('Objective and domain scope','different owners / earlier stage / fixed coordinate','scope retained; unavailable curvature incomplete')
]

edge_rows += [
 ('Retained payoff ceilings','A=2^53 / 2^60 / 10^20; fixed decisions','three excesses fail; equal/below pass; Bounded tolerance retained'),
 ('Strict coefficient signs','cost=-10^-400; parameters / mutations / expressions','negative rejects before search; zero/positive controls pass; underflow range explicit')
]

edge_rows += [
 ('Mixed endpoint retention','Float 1 / Rational 1 plus or minus 10^-20','reversed rejects; collapsed positive width incomplete; ordinary controls pass'),
 ('Tolerance admission','negative / positive 10^-400; retained rational','negative and range errors; exact zero and subnormal controls pass'),
 ('Finite-real payoff admission','net or gross with imaginary 10^-400','known invalid fails; unknown incomplete; legal negative net passes')
]

edge_rows += [
 ('Owner payoff admission','payoff_for; fixed / interior / lower / skipped','known invalid fails; unknown incomplete; finite negative controls pass'),
 ('Retained payoff differences','A=2^53 / 2^60 / 10^20; exact gap 1','RL ranks first; weak reverse fails; crossing .5; slope 1')
]

def font(run,name='Times New Roman',size=11):
 run.font.name=name;run.font.size=Pt(size);run.font.color.rgb=RGBColor(0,0,0)
 run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
def doc_base():
 d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
 s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.left_margin=Cm(2.25);s.right_margin=Cm(2.25)
 styles=d.styles
 for n in ('Normal','Body Text','Caption','Title','Heading 1','Heading 2','Heading 3'):
  st=styles[n];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(11)
  st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
 styles['Normal'].paragraph_format.space_after=Pt(6);styles['Normal'].paragraph_format.line_spacing=1.06
 styles['Heading 1'].font.size=Pt(13);styles['Heading 1'].font.bold=True
 styles['Heading 2'].font.size=Pt(11);styles['Heading 2'].font.bold=True
 styles['Title'].font.size=Pt(17);styles['Title'].font.bold=True
 for border in styles['Title']._element.findall('.//'+qn('w:pBdr')):
  border.getparent().remove(border)
 f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
 z=OxmlElement('w:fldSimple');z.set(qn('w:instr'),'PAGE');f._p.append(z)
 return d
def para(d,text,style=None):
 p=d.add_paragraph(style=style);run=p.add_run(text);font(run, size=10 if style=='Caption' else 11)
 if style is None:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 return p
def table(d,rows,widths=None):
 t=d.add_table(rows=0,cols=len(rows[0]));t.style='Table Grid';t.autofit=False
 if widths:
  for column,width in zip(t.columns,widths):column.width=Cm(width)
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
 t._tbl.tblPr.append(borders)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for j,v in enumerate(row):
   cells[j].text=str(v)
   cells[j].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   margins=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    m=OxmlElement('w:'+edge);m.set(qn('w:w'),'80');m.set(qn('w:type'),'dxa');margins.append(m)
   cells[j]._tc.get_or_add_tcPr().append(margins)
   for p in cells[j].paragraphs:
    p.paragraph_format.space_after=Pt(3)
    for run in p.runs:font(run,size=9);run.bold=(i==0)
   if i==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');cells[j]._tc.get_or_add_tcPr().append(sh)
   if widths:cells[j].width=Cm(widths[j])
  trpr=t.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if i==0:
   repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
 d.add_paragraph()
 return t

d=doc_base();d.add_heading(c['title'],0);para(d,c['author_note'])
para(d,'Original Software Publication • Review revision 0.2.11 • 8 October 2026')
d.add_heading('Highlights',1)
for x in c['highlights']:para(d,'• '+x)
d.add_heading('Abstract',1);para(d,c['abstract']);para(d,'Keywords: '+c['keywords'])
para(d,'Table 1. Code metadata.', 'Caption');table(d,[('Nr.','Description','Value')]+metadata1,[1,4.4,11.1])
para(d,'Table 2. Software metadata.', 'Caption');table(d,[('Nr.','Description','Value')]+metadata2,[1,4.4,11.1])

def render_word(items):
 for x in items:
  if isinstance(x,str):para(d,x)
  elif x[0]=='code':
   p=para(d,x[1]);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
   for run in p.runs:font(run,'Consolas',8)
  else:d.add_heading(x[0],2);render_word(x[1])
for title,items in c['sections']:
 d.add_heading(title,1);render_word(items)
 if title.startswith('1.'):
  para(d,'Table 3. Scope comparison based on the tools’ own descriptions.', 'Caption');table(d,comparison,[2.1,4.8,9.6])
 if title.startswith('2.'):
  d.add_picture(str(figdir/'fig1_architecture.png'),width=Cm(16.3));para(d,figures[0][1],'Caption')
 if title.startswith('3.'):
  para(d,'Table 4. Illustrative numerical candidates (sigma=0.1, phi=0.85, growth_end=0.3; kappa_d=0.05 in variants).', 'Caption');table(d,numerics,[3.1,3.1,2.2,4.2,3.9])
  for tag,caption in figures[1:]:d.add_picture(str(figdir/(tag+'.png')),width=Cm(16.3));para(d,caption,'Caption')
  para(d,'Table 5. Selected boundaries from the eleven numbered audit acceptance files 021 through 029, 0210 and 0211. The last two rows summarize the R12 fixes and controls.', 'Caption')
  table(d,edge_rows,[3.9,5.2,7.4])
for title,text in c['declarations']:
 d.add_heading(title,1);para(d,text).paragraph_format.keep_together=True
d.add_heading('References',1)
for ref in c['references']:
 p=para(d,ref);p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_together=True
 for run in p.runs:font(run,size=10)
path=root/'PACEgame_SoftwareX_Manuscript_v0_2_11.docx';d.save(path);shutil.copy2(path,pkg/path.name)
# One content source produces the editable Word manuscript and its TeX proof.
def esc(s):
 maps={'^':r'\textasciicircum{}','~':r'\textasciitilde{}','\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','×':r'$\times$','≤':r'$\leq$','−':r'$-$'}
 s=s.replace('⁻¹⁶','\ue000').replace('⁻⁶','\ue001').replace('⁻¹⁷','\ue002');maps.update({'\ue000':r'$^{-16}$','\ue001':r'$^{-6}$','\ue002':r'$^{-17}$'})
 return ''.join(maps.get(c,c) for c in s)

tex=[r'\documentclass[11pt,a4paper]{article}',r'\usepackage[margin=2.1cm]{geometry}',r'\usepackage{fontspec}',r'\IfFontExistsTF{TeX Gyre Termes}{\setmainfont{TeX Gyre Termes}}{\setmainfont{Latin Modern Roman}}',r'\setmonofont{DejaVu Sans Mono}',r'\usepackage{graphicx,longtable,booktabs,array,hyperref,xurl,needspace,placeins}',r'\hypersetup{hidelinks}',r'\setlength{\parindent}{0pt}',r'\setlength{\parskip}{4pt}',r'\begin{document}',r'\begin{center}\Large\bfseries '+esc(c['title'])+r'\end{center}',esc(c['author_note'])+r'\par',r'\textit{Original Software Publication; review revision 0.2.11; 8 October 2026.}',r'\subsection*{Highlights}',r'\begin{itemize}']
tex += [r'\item '+esc(x) for x in c['highlights']];tex +=[r'\end{itemize}',r'\subsection*{Abstract}',esc(c['abstract']),r'\par\textbf{Keywords:} '+esc(c['keywords'])]
def tex_table(rows,caption,cols):
 tex.extend([r'\begingroup\small',r'\begin{longtable}{'+cols+'}',r'\caption{'+esc(caption)+r'}\\',r'\toprule', ' & '.join(r'\textbf{'+esc(x)+'}' for x in rows[0])+r'\\\midrule\endfirsthead',r'\toprule',' & '.join(r'\textbf{'+esc(x)+'}' for x in rows[0])+r'\\\midrule\endhead'])
 for row in rows[1:]:tex.append(' & '.join(esc(str(x)) for x in row)+r'\\')
 tex.extend([r'\bottomrule\end{longtable}\endgroup'])
tex_table([('Nr.','Description','Value')]+metadata1,'Code metadata.',r'p{.6cm}p{3.9cm}p{10.7cm}')
tex_table([('Nr.','Description','Value')]+metadata2,'Software metadata.',r'p{.6cm}p{3.9cm}p{10.7cm}')
def render_tex(items):
 for x in items:
  if isinstance(x,str):tex.append(esc(x)+r'\par')
  elif x[0]=='code':tex.extend([r'{\small\begin{verbatim}',x[1],r'\end{verbatim}}'])
  else:tex.append(r'\subsection*{'+esc(x[0])+'}');render_tex(x[1])
def figure(tag,caption):
 tex.extend([r'\begin{figure}[htbp]\centering',r'\includegraphics[width=\linewidth]{figures/'+tag+'.pdf}',r'\caption{'+esc(re.sub(r'^Fig\. \d+\. ', '', caption))+r'}\end{figure}'])
for title,items in c['sections']:
 tex.append(r'\section*{'+esc(title)+'}');render_tex(items)
 if title.startswith('1.'):tex_table(comparison,'Scope comparison.',r'p{2cm}p{4cm}p{9cm}')
 if title.startswith('2.'):
  figure(*figures[0]);tex.append(r'\FloatBarrier')
 if title.startswith('3.'):
  tex_table(numerics,'Illustrative numerical candidates.',r'p{2.6cm}p{2.4cm}p{1.7cm}p{4cm}p{3.3cm}')
  for args in figures[1:]:figure(*args)
  tex_table(edge_rows,'Selected boundary acceptance results.',r'p{3.1cm}p{5.1cm}p{7cm}')
for title,text in c['declarations']:tex.append(r'\section*{'+esc(title)+'}');tex.append(esc(text))
tex.append(r'\Needspace{6cm}');tex.append(r'\section*{References}');tex.append(r'\begingroup\small')
for ref in c['references']:
 parts=re.split(r'(https?://\S+)',ref)
 tex.append(''.join(r'\url{'+part.rstrip('.')+'}' if part.startswith('http') else esc(part) for part in parts)+r'\par')
tex.append(r'\endgroup');tex.append(r'\end{document}')
(pkg/'manuscript.tex').write_text('\n'.join(tex))
(pkg/'manuscript_content.json').write_text(json.dumps(c,ensure_ascii=False,indent=2))
(pkg/'README.txt').write_text(f'''Complete Word manuscript and XeLaTeX proof, generated from the same content source. Body word count: {c['body_word_count']}; abstract: {c['abstract_word_count']}.\nCompile manuscript.tex with xelatex twice. TeX Gyre Termes or Latin Modern Roman, DejaVu Sans Mono and standard TeX Live packages are required.\nThis is an OSP-structured review manuscript. The current official template file could not be retrieved from Elsevier; final official-template verification remains pending. No claim of official-template certification is made.\nAuthor identity, support contact, public code/archive identifiers and author declarations remain to be confirmed.\n''')
(pkg/'metadata_tables.json').write_text(json.dumps({'code':metadata1,'software':metadata2},indent=2))
print(path)
