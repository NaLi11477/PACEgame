from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2);sec.right_margin=Cm(2);sec.footer_distance=Cm(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=d.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PACEgame Revision 0211 CJK')
 for b in st.element.findall('.//'+qn('w:pBdr')):b.getparent().remove(b)
d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.line_spacing=Pt(14.5);d.styles['Normal'].paragraph_format.space_after=Pt(5)
d.styles['Title'].font.size=Pt(19);d.styles['Title'].paragraph_format.line_spacing=Pt(25);d.styles['Title'].paragraph_format.space_after=Pt(10)
for name,size,line in [('Heading 1',13,18),('Heading 2',11.5,16)]:
 st=d.styles[name];st.font.size=Pt(size);st.paragraph_format.line_spacing=Pt(line);st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(6)
def p(s):return d.add_paragraph(s)
def lead(label,s):
 x=d.add_paragraph();x.add_run(label+'：').bold=True;x.add_run(s);return x
def h(s,newpage=False):
 x=d.add_heading(s,1)
 if newpage:x.paragraph_format.page_break_before=True
 return x
def table(rows,widths,center=()):
 t=d.add_table(rows=0,cols=len(widths));t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
 t._tbl.tblPr.append(borders)
 for i,row in enumerate(rows):
  cells=t.add_row().cells;pr=t.rows[-1]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for j,value in enumerate(row):
   cell=cells[j];cell.width=Cm(widths[j]);cell.text=str(value);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   margins=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:w'),'85');x.set(qn('w:type'),'dxa');margins.append(x)
   cell._tc.get_or_add_tcPr().append(margins)
   if i==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDEDED');cell._tc.get_or_add_tcPr().append(shade)
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(2);par.paragraph_format.space_before=Pt(2);par.paragraph_format.line_spacing=Pt(13.5)
    if j in center:par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in par.runs:run.font.size=Pt(10);run.bold=i==0
 d.add_paragraph().paragraph_format.space_after=Pt(0)
 return t
def small(s):
 x=p(s)
 for r in x.runs:r.font.size=Pt(9)
 return x
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

